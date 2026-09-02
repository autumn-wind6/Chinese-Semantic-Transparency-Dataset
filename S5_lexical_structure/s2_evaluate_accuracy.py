#!/usr/bin/env python3
"""Evaluate the 12-class human-curated lexical-structure benchmark."""

from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def evaluate(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    summary_rows = []
    error_rows = []
    for sheet in workbook.worksheets:
        expected = sheet.title.strip()
        total = correct = filled = 0
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            word = "" if row[0] is None else str(row[0]).strip()
            if not word:
                continue
            actual = "" if len(row) < 2 or row[1] is None else str(row[1]).strip()
            total += 1
            filled += bool(actual)
            correct += actual == expected
            if actual != expected:
                error_rows.append({"sheet": sheet.title, "row": row_number, "word": word, "expected": expected, "predicted": actual})
        summary_rows.append(
            {
                "structure": expected,
                "n": total,
                "filled": filled,
                "correct": correct,
                "errors": total - correct,
                "accuracy": correct / total if total else float("nan"),
            }
        )
    workbook.close()
    summary = pd.DataFrame(summary_rows)
    summary.loc[len(summary)] = {
        "structure": "Overall",
        "n": int(summary["n"].sum()),
        "filled": int(summary["filled"].sum()),
        "correct": int(summary["correct"].sum()),
        "errors": int(summary["errors"].sum()),
        "accuracy": summary["correct"].sum() / summary["n"].sum(),
    }
    return summary, pd.DataFrame(error_rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--errors", type=Path, required=True)
    args = parser.parse_args()
    summary, errors = evaluate(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.errors.parent.mkdir(parents=True, exist_ok=True)
    summary.to_csv(args.output, index=False)
    errors.to_csv(args.errors, index=False)
    print(summary.to_string(index=False))


if __name__ == "__main__":
    main()

